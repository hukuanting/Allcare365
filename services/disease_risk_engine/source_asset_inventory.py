"""Read-only inventory for the professor's disease-risk source archive.

This module deliberately extracts evidence, not executable clinical logic.
Excel formula cells, Office document structure, citations, hashes, and review
status are recorded so that a reviewer can promote one algorithm at a time.
Ambiguous or legacy formats remain explicitly gated.
"""

from __future__ import annotations

import ast
import hashlib
import json
import re
import zipfile
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree

from openpyxl import load_workbook

from .algorithm_registry import RUNTIME_ALGORITHMS


INVENTORY_VERSION = "1.0.0"
FORMULA_SNIPPET_LIMIT = 100
FORMULA_SNIPPET_LENGTH = 400

_WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_PRESENTATION_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
_DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

_FORMULA_TEXT = re.compile(
    r"(?i)(?:=|\b(?:formula|equation|coefficient|score|risk|odds|log|ln|exp|sqrt)\b|β|×|÷)"
)
_URL = re.compile(r"https?://[^\s<>'\"]+")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _clean_text(value: str) -> str:
    return " ".join(str(value or "").split())


def _formula_snippets(values: Iterable[str]) -> list[str]:
    snippets = []
    seen = set()
    for value in values:
        normalized = _clean_text(value)
        if not normalized or not _FORMULA_TEXT.search(normalized):
            continue
        normalized = normalized[:FORMULA_SNIPPET_LENGTH]
        if normalized in seen:
            continue
        seen.add(normalized)
        snippets.append(normalized)
        if len(snippets) >= FORMULA_SNIPPET_LIMIT:
            break
    return snippets


def _xlsx_metadata(path: Path) -> dict[str, Any]:
    workbook = load_workbook(
        path,
        read_only=True,
        data_only=False,
        keep_links=False,
    )
    sheets = []
    total_formulas = 0
    formulas = []
    try:
        for worksheet in workbook.worksheets:
            sheet_formulas = []
            for row in worksheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if cell.data_type == "f" or (
                        isinstance(value, str) and value.startswith("=")
                    ):
                        expression = str(value)
                        sheet_formulas.append(
                            {"cell": cell.coordinate, "expression": expression}
                        )
                        formulas.append(
                            {
                                "sheet": worksheet.title,
                                "cell": cell.coordinate,
                                "expression": expression,
                            }
                        )
            total_formulas += len(sheet_formulas)
            sheets.append(
                {
                    "name": worksheet.title,
                    "state": worksheet.sheet_state,
                    "max_row": worksheet.max_row,
                    "max_column": worksheet.max_column,
                    "formula_count": len(sheet_formulas),
                }
            )
    finally:
        workbook.close()
    with zipfile.ZipFile(path) as archive:
        external_link_parts = sum(
            name.startswith("xl/externalLinks/externalLink") and name.endswith(".xml")
            for name in archive.namelist()
        )
    return {
        "extraction_status": (
            "formula_cells_extracted" if total_formulas else "no_formula_cells_detected"
        ),
        "sheet_count": len(sheets),
        "sheets": sheets,
        "formula_count": total_formulas,
        "formulas": formulas,
        "external_link_parts": external_link_parts,
    }


def _xml_text(root: ElementTree.Element, text_tag: str) -> list[str]:
    values = []
    for element in root.iter(text_tag):
        if element.text:
            values.append(element.text)
    return values


def _external_targets(archive: zipfile.ZipFile, relationship_parts: Iterable[str]) -> list[str]:
    targets = set()
    for part in relationship_parts:
        if part not in archive.namelist():
            continue
        root = ElementTree.fromstring(archive.read(part))
        for relationship in root.iter(f"{{{_REL_NS}}}Relationship"):
            if relationship.attrib.get("TargetMode") != "External":
                continue
            target = relationship.attrib.get("Target", "").strip()
            if target:
                targets.add(target)
    return sorted(targets)


def _docx_metadata(path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
        paragraphs = []
        for paragraph in root.iter(f"{{{_WORD_NS}}}p"):
            text = _clean_text("".join(_xml_text(paragraph, f"{{{_WORD_NS}}}t")))
            if text:
                paragraphs.append(text)
        targets = _external_targets(
            archive,
            ["word/_rels/document.xml.rels"],
        )
        embedded_objects = sum(name.startswith("word/embeddings/") for name in archive.namelist())
        images = sum(name.startswith("word/media/") for name in archive.namelist())
        math_objects = sum(1 for _ in root.iter(f"{{{_MATH_NS}}}oMath"))
        table_count = sum(1 for _ in root.iter(f"{{{_WORD_NS}}}tbl"))
    return {
        "extraction_status": "text_indexed_requires_formula_review",
        "paragraph_count": len(paragraphs),
        "table_count": table_count,
        "math_object_count": math_objects,
        "embedded_object_count": embedded_objects,
        "image_count": images,
        "formula_candidate_snippets": _formula_snippets(paragraphs),
        "external_targets": targets,
        "detected_urls": sorted({url.rstrip(".,);") for text in paragraphs for url in _URL.findall(text)}),
    }


def _pptx_metadata(path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(path) as archive:
        slide_parts = sorted(
            name
            for name in archive.namelist()
            if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
        )
        slide_text = []
        for part in slide_parts:
            root = ElementTree.fromstring(archive.read(part))
            text = _clean_text("".join(_xml_text(root, f"{{{_DRAWING_NS}}}t")))
            if text:
                slide_text.append(text)
        relationship_parts = [
            name
            for name in archive.namelist()
            if name.startswith("ppt/slides/_rels/") and name.endswith(".rels")
        ]
        targets = _external_targets(archive, relationship_parts)
        images = sum(name.startswith("ppt/media/") for name in archive.namelist())
    return {
        "extraction_status": "slides_indexed_requires_formula_review",
        "slide_count": len(slide_parts),
        "image_count": images,
        "formula_candidate_snippets": _formula_snippets(slide_text),
        "external_targets": targets,
    }


def _python_metadata(path: Path) -> dict[str, Any]:
    source = path.read_text(encoding="utf-8")
    tree = ast.parse(source, filename=str(path))
    functions = [
        node.name
        for node in ast.walk(tree)
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))
    ]
    return {
        "extraction_status": "derived_python_catalog_indexed",
        "function_count": len(functions),
        "function_names": sorted(functions),
    }


def inspect_asset(path: Path, source_root: Path) -> dict[str, Any]:
    extension = path.suffix.lower()
    role = "source_evidence"
    metadata: dict[str, Any]
    try:
        if extension == ".xlsx":
            metadata = _xlsx_metadata(path)
        elif extension == ".docx":
            metadata = _docx_metadata(path)
        elif extension == ".pptx":
            metadata = _pptx_metadata(path)
        elif extension == ".py":
            role = "derived_code"
            metadata = _python_metadata(path)
        elif extension == ".md":
            role = "review_report"
            text = path.read_text(encoding="utf-8")
            metadata = {
                "extraction_status": "review_report_indexed",
                "line_count": len(text.splitlines()),
                "detected_urls": sorted({url.rstrip(".,);") for url in _URL.findall(text)}),
            }
        elif extension == ".xls":
            metadata = {"extraction_status": "legacy_excel_requires_lossless_conversion"}
        elif extension == ".doc":
            metadata = {"extraction_status": "legacy_word_requires_lossless_conversion"}
        elif extension == ".pdf":
            metadata = {"extraction_status": "pdf_requires_text_and_visual_review"}
        elif extension == ".png":
            metadata = {"extraction_status": "image_requires_visual_review"}
        else:
            metadata = {"extraction_status": "unsupported_or_empty_asset"}
    except Exception as exc:  # An inventory must record, not hide, unreadable evidence.
        metadata = {
            "extraction_status": "extraction_error",
            "error_type": type(exc).__name__,
            "error": str(exc)[:500],
        }

    stat = path.stat()
    return {
        "relative_path": path.relative_to(source_root).as_posix(),
        "file_name": path.name,
        "extension": extension,
        "size_bytes": stat.st_size,
        "modified_at": datetime.fromtimestamp(
            stat.st_mtime,
            tz=timezone.utc,
        ).isoformat(),
        "sha256": _sha256(path),
        "source_role": role,
        **metadata,
    }


def _runtime_catalog() -> list[dict[str, Any]]:
    return [
        {
            "algorithm_id": item.algorithm_id,
            "display_name": item.display_name,
            "model_version": item.model_version,
            "source_label": item.source_label,
            "method_uri": item.method_uri,
            "time_horizon": item.time_horizon,
            "required_inputs": list(item.required_inputs),
            "fhir_output": item.output_resource.value,
            "governance_status": item.governance_status,
            "runtime_enabled": item.runtime_enabled,
        }
        for item in RUNTIME_ALGORITHMS
    ]


def build_source_manifest(source_root: Path | str) -> dict[str, Any]:
    root = Path(source_root).resolve()
    if not root.is_dir():
        raise ValueError(f"Source archive directory does not exist: {root}")
    assets = [
        inspect_asset(path, root)
        for path in sorted(root.rglob("*"), key=lambda item: item.as_posix().casefold())
        if path.is_file()
    ]
    hashes: dict[str, list[str]] = defaultdict(list)
    for asset in assets:
        hashes[asset["sha256"]].append(asset["relative_path"])
    duplicate_groups = [
        {"sha256": digest, "files": paths}
        for digest, paths in sorted(hashes.items())
        if len(paths) > 1
    ]
    runtime = _runtime_catalog()
    status_counts: dict[str, int] = defaultdict(int)
    for asset in assets:
        status_counts[asset["extraction_status"]] += 1
    return {
        "manifest_version": INVENTORY_VERSION,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source_root_name": root.name,
        "archive_sha256": hashlib.sha256(
            "\n".join(
                f"{asset['relative_path']}\t{asset['sha256']}" for asset in assets
            ).encode("utf-8")
        ).hexdigest(),
        "summary": {
            "asset_count": len(assets),
            "total_bytes": sum(asset["size_bytes"] for asset in assets),
            "xlsx_formula_cells": sum(asset.get("formula_count", 0) for asset in assets),
            "runtime_algorithm_count": len(runtime),
            "runtime_missing_method_uri": sum(not item["method_uri"] for item in runtime),
            "extraction_status_counts": dict(sorted(status_counts.items())),
        },
        "governance": {
            "clinical_execution_allowed": False,
            "promotion_rule": (
                "Extraction is evidence only. A model requires primary-source verification, "
                "typed inputs and units, applicability rules, independent test vectors, "
                "clinical-owner approval, and runtime registry promotion."
            ),
        },
        "duplicate_groups": duplicate_groups,
        "runtime_catalog": runtime,
        "assets": assets,
    }


def write_source_manifest(source_root: Path | str, output_path: Path | str) -> dict[str, Any]:
    manifest = build_source_manifest(source_root)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return manifest


__all__ = [
    "INVENTORY_VERSION",
    "build_source_manifest",
    "inspect_asset",
    "write_source_manifest",
]
